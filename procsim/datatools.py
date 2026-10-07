"""Data exchange (v7.4): import unit parameters, feed compositions, profile tables and dynamic events from CSV / Excel /
YAML, batch edits, result tables, export to Excel / CSV / YAML, and a restricted Python runner for adjusting the tables
before they are exported.  No Streamlit dependency.

Values are always in the model's own units (the catalogue unit shown next to every parameter: bar(a), °C, kW, ...),
whatever the display-unit system of the app.

Import works in two steps so that nothing changes silently: ``requests_from_*`` read a file into requests,
``plan`` checks them against the flowsheet and returns one row per change (old value, new value, status), and
``apply_plan`` writes the accepted rows into the model.
"""
from __future__ import annotations

import ast
import copy
import difflib
import fnmatch
import io
import json
import math
import os
import re
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import OrderedDict

import numpy as np
import pandas as pd

from .unitops import CATALOGUE, default_params

SEP = " | "
COMP_PREFIX = "composition."
UNIT_COLS = ("unit", "name", "block", "tag", "unit name", "object", "feed", "stream", "units")
PARAM_COLS = ("parameter", "param", "key", "property", "variable", "setting", "parameter key")
VALUE_COLS = ("value", "new value", "val", "new")
COMP_COLS = ("component", "comp", "species")
FRAC_COLS = ("fraction", "value", "frac", "mole fraction", "x")
MAX_CHANGES = 20000
STATUS_ORDER = ("error", "warning", "ok", "unchanged", "skipped")


class DataError(ValueError):
    pass


# ================================================================================================== model lookups
def _norm(s):
    return re.sub(r"\s+", " ", str(s).strip().lower())


def unit_names(model):
    """({name: uid}, {names used by more than one unit})."""
    by, dup = {}, set()
    for uid, u in model["units"].items():
        if u["name"] in by:
            dup.add(u["name"])
        by[u["name"]] = uid
    return by, dup


def components(model):
    fl = model["fluid"]
    return list(fl["components"])


def param_specs(utype):
    return {p["key"]: p for p in CATALOGUE[utype]["params"]}


def _is_blank(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str) and not v.strip():
        return True
    try:
        return bool(pd.isna(v)) if not isinstance(v, (list, dict, str)) else False
    except (TypeError, ValueError):
        return False


def to_number(v):
    """float from a cell (accepts a decimal comma and thin / normal spaces); None when blank; ValueError otherwise."""
    if _is_blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError("a yes / no value is not a number")
    if isinstance(v, (int, float, np.integer, np.floating)):
        x = float(v)
    else:
        s = str(v).strip().replace(" ", "").replace(" ", "").replace(" ", "")
        if s.count(",") == 1 and "." not in s:
            s = s.replace(",", ".")
        elif s.count(",") >= 1 and "." in s:
            s = s.replace(",", "")
        x = float(s)
    if not math.isfinite(x):
        raise ValueError("not a finite number")
    return x


def strip_unit(h):
    """'Pressure [bar(a)]' -> 'Pressure'."""
    h = str(h).strip()
    if h.endswith("]") and " [" in h:
        h = h[: h.rindex(" [")]
    return h.strip()


# ================================================================================================== current values as tables
def parameter_table(model):
    """Long table of every catalogue parameter of every unit: Unit, Type, Parameter (key), Label, Value, Unit of measure,
    Options.  Written as the 'Parameters' sheet of an export and read back by the import."""
    rows = []
    for uid, u in model["units"].items():
        cat = CATALOGUE[u["type"]]
        for p in cat["params"]:
            v = u["params"].get(p["key"], p.get("default"))
            rows.append({"Unit": u["name"], "Type": cat["label"], "Parameter": p["key"], "Label": p["label"],
                         "Value": v, "Unit of measure": p.get("unit", "") or "",
                         "Options": " | ".join(map(str, p["options"])) if p.get("kind") == "select" else ""})
    return pd.DataFrame(rows, columns=["Unit", "Type", "Parameter", "Label", "Value", "Unit of measure", "Options"])


def composition_table(model):
    """Long table Feed, Component, Fraction for every feed (as entered: mole or mass fractions per its Composition basis)."""
    comps = components(model)
    rows = []
    for u in model["units"].values():
        if u["type"] != "feed":
            continue
        c = u["params"].get("composition") or {}
        for k in comps:
            rows.append({"Feed": u["name"], "Component": k, "Fraction": float(c.get(k, 0.0) or 0.0),
                         "Basis": u["params"].get("comp_basis", "Mole fractions")})
    return pd.DataFrame(rows, columns=["Feed", "Component", "Fraction", "Basis"])


def composition_wide(model):
    """One row per feed, one column per component."""
    comps = components(model)
    rows = []
    for u in model["units"].values():
        if u["type"] != "feed":
            continue
        c = u["params"].get("composition") or {}
        r = {"Feed": u["name"]}
        r.update({k: float(c.get(k, 0.0) or 0.0) for k in comps})
        rows.append(r)
    return pd.DataFrame(rows, columns=["Feed"] + comps)


def wide_parameters(model, utype):
    """One row per unit of a type, one column per catalogue parameter (for the table editor)."""
    cat = CATALOGUE[utype]
    rows = []
    for u in model["units"].values():
        if u["type"] != utype:
            continue
        r = {"Unit": u["name"]}
        for p in cat["params"]:
            r[p["key"]] = u["params"].get(p["key"], p.get("default"))
        rows.append(r)
    return pd.DataFrame(rows, columns=["Unit"] + [p["key"] for p in cat["params"]])


def events_table(model):
    from . import dynamics as DY
    rows = []
    for e in (model.get("dynamics") or {}).get("events") or []:
        rows.append({"Time [s]": e.get("t"), "Event": DY.EVENT_KINDS.get(e.get("kind"), (e.get("kind"),))[0],
                     "Target": e.get("target"), "Value": e.get("value"), "Ramp [s]": e.get("ramp", 0.0)})
    return pd.DataFrame(rows, columns=["Time [s]", "Event", "Target", "Value", "Ramp [s]"])


# ================================================================================================== reading files
def _decode(data):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def read_tables(filename, data):
    """{sheet name: DataFrame of strings / numbers as in the file} from CSV, TSV, Excel."""
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=0, engine="openpyxl")
        except Exception as e:
            raise DataError(f"could not read the Excel file: {e}")
        out = OrderedDict()
        for nm, df in sheets.items():
            df = df.dropna(how="all").dropna(axis=1, how="all")
            df.columns = [str(c).strip() for c in df.columns]
            if len(df.columns):
                out[str(nm)] = df.reset_index(drop=True)
        if not out:
            raise DataError("the workbook has no data")
        return out
    if ext in (".csv", ".tsv", ".txt"):
        text = _decode(data)
        try:
            df = pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=object, skip_blank_lines=True)
        except Exception as e:
            raise DataError(f"could not read the CSV file: {e}")
        df = df.dropna(how="all").dropna(axis=1, how="all")
        df.columns = [str(c).strip() for c in df.columns]
        if df.empty:
            raise DataError("the file has no data rows")
        return OrderedDict([(os.path.splitext(os.path.basename(filename))[0] or "csv", df.reset_index(drop=True))])
    raise DataError(f"unsupported file type '{ext}' (use .csv, .xlsx or .yaml)")


# ================================================================================================== requests
def _req(unit, param, value, src="", row=None):
    return {"unit": unit, "param": param, "value": value, "src": src, "row": row}


def _find_col(cols, aliases):
    low = {_norm(c): c for c in cols}
    for a in aliases:
        if a in low:
            return low[a]
    low2 = {_norm(strip_unit(c)): c for c in cols}
    for a in aliases:
        if a in low2:
            return low2[a]
    return None


KINDS = ("Parameters (one row per parameter)", "Parameters (one row per unit)", "Feed compositions (one row per component)",
         "Feed compositions (one row per feed)", "Profile table", "Dynamic events")


_PARAM_NAMES = None


def _all_param_names():
    global _PARAM_NAMES
    if _PARAM_NAMES is None:
        names = set()
        for c in CATALOGUE.values():
            for p in c["params"]:
                names.add(_norm(p["key"]))
                names.add(_norm(p["label"]))
                names.add(_norm(strip_unit(p["label"])))
        _PARAM_NAMES = names
    return _PARAM_NAMES


def classify(df, model):
    """(kind, reason) of a table, or (None, reason) when it cannot be recognised."""
    cols = list(df.columns)
    if not cols:
        return None, "no columns"
    first = _norm(strip_unit(cols[0]))
    ucol = _find_col(cols, UNIT_COLS)
    if first in ("time", "t", "step") and any(SEP in c for c in cols[1:]):
        return KINDS[4], "first column is Time and the others are 'Unit | parameter'"
    if _find_col(cols, ("event",)) and _find_col(cols, ("target",)) and _find_col(cols, ("time", "t")):
        return KINDS[5], "has Time, Event and Target columns"
    if ucol and _find_col(cols, COMP_COLS) and _find_col(cols, FRAC_COLS):
        return KINDS[2], "has Feed, Component and Fraction columns"
    if ucol and _find_col(cols, PARAM_COLS) and _find_col(cols, VALUE_COLS):
        return KINDS[0], "has Unit, Parameter and Value columns"
    if ucol and ucol == cols[0]:
        rest = [strip_unit(c) for c in cols if c != ucol]
        rest = [c for c in rest if _norm(c) not in ("sum", "total", "type", "basis", "status", "label", "value", "quantity", "message")]
        comps = set(components(model))
        if rest and all(c in comps for c in rest):
            return KINDS[3], "the columns are components"
        known = _all_param_names()
        hits = [c for c in rest if _norm(c) in known or c.startswith(COMP_PREFIX)]
        if hits:
            return KINDS[1], "first column is the unit name, the others are parameters"
        return None, "the columns are not parameter names of any unit type (a results sheet?)"
    return None, "no Unit / Name column (first column should be the unit name)"


def requests_from_frame(df, model, kind=None):
    """(requests, extras, notes) from one table.  extras carries 'profile_rows' / 'events' for those kinds."""
    reason = ""
    if kind is None:
        kind, reason = classify(df, model)
    if kind is None:
        raise DataError(f"table not recognised: {reason}")
    cols = list(df.columns)
    reqs, extras, notes = [], {}, []
    if kind == KINDS[0]:
        uc, pc, vc = _find_col(cols, UNIT_COLS), _find_col(cols, PARAM_COLS), _find_col(cols, VALUE_COLS)
        if not (uc and pc and vc):
            raise DataError("needs the columns Unit, Parameter and Value")
        for i, r in df.iterrows():
            if _is_blank(r[uc]) or _is_blank(r[pc]):
                continue
            reqs.append(_req(str(r[uc]).strip(), str(r[pc]).strip(), r[vc], "row", i + 2))
    elif kind == KINDS[1]:
        uc = _find_col(cols, UNIT_COLS)
        if not uc:
            raise DataError("the first column must be the unit name (header 'Unit' or 'Name')")
        skip = {"type", "label", "status", "sum", "total", "basis"}
        for i, r in df.iterrows():
            if _is_blank(r[uc]):
                continue
            for c in cols:
                if c == uc or _norm(strip_unit(c)) in skip:
                    continue
                if _is_blank(r[c]):
                    continue
                reqs.append(_req(str(r[uc]).strip(), strip_unit(c), r[c], "row", i + 2))
    elif kind == KINDS[2]:
        uc, cc, fc = _find_col(cols, UNIT_COLS), _find_col(cols, COMP_COLS), _find_col(cols, FRAC_COLS)
        if not (uc and cc and fc):
            raise DataError("needs the columns Feed, Component and Fraction")
        for i, r in df.iterrows():
            if _is_blank(r[uc]) or _is_blank(r[cc]):
                continue
            reqs.append(_req(str(r[uc]).strip(), COMP_PREFIX + str(r[cc]).strip(), r[fc], "row", i + 2))
    elif kind == KINDS[3]:
        uc = _find_col(cols, UNIT_COLS)
        for i, r in df.iterrows():
            if _is_blank(r[uc]):
                continue
            for c in cols:
                if c == uc or _norm(strip_unit(c)) in ("sum", "total", "basis"):
                    continue
                if _is_blank(r[c]):
                    continue
                reqs.append(_req(str(r[uc]).strip(), COMP_PREFIX + strip_unit(c), r[c], "row", i + 2))
    elif kind == KINDS[4]:
        rows = []
        head = ["Time"] + [strip_unit(c) for c in cols[1:]]
        for i, r in df.iterrows():
            row = {}
            for h, c in zip(head, cols):
                try:
                    row[h] = to_number(r[c])
                except ValueError:
                    raise DataError(f"row {i + 2}, column '{c}': '{r[c]}' is not a number")
            if any(v is not None for v in row.values()):
                rows.append(row)
        if not rows:
            raise DataError("the profile table has no rows")
        extras["profile_rows"] = rows
    elif kind == KINDS[5]:
        from . import dynamics as DY
        tc, ec, gc = _find_col(cols, ("time", "t")), _find_col(cols, ("event",)), _find_col(cols, ("target",))
        vc, rc = _find_col(cols, ("value",)), _find_col(cols, ("ramp",))
        label2kind = {v[0].lower(): k for k, v in DY.EVENT_KINDS.items()}
        evs = []
        for i, r in df.iterrows():
            if _is_blank(r[ec]):
                continue
            k = str(r[ec]).strip()
            kind_ = k if k in DY.EVENT_KINDS else label2kind.get(k.lower())
            if kind_ is None:
                raise DataError(f"row {i + 2}: unknown event '{k}' (known: {', '.join(v[0] for v in DY.EVENT_KINDS.values())})")
            try:
                t = to_number(r[tc])
            except ValueError:
                raise DataError(f"row {i + 2}: the time is not a number")
            if t is None or t < 0:
                raise DataError(f"row {i + 2}: the time must be a number of seconds, zero or later")
            e = {"t": t, "kind": kind_, "target": None if _is_blank(r[gc]) else str(r[gc]).strip()}
            if vc is not None and not _is_blank(r[vc]):
                try:
                    e["value"] = to_number(r[vc])
                except ValueError:
                    e["value"] = str(r[vc]).strip()           # e.g. Auto / Manual
            if rc is not None and not _is_blank(r[rc]):
                e["ramp"] = to_number(r[rc])
            evs.append(e)
        extras["events"] = evs
    notes.append(f"{kind}" + (f" ({reason})" if reason else ""))
    return reqs, extras, notes


def requests_from_yaml(text, model):
    """(requests, extras, notes) from YAML.  A whole flowsheet (fluid + units + streams) comes back as extras['flowsheet']."""
    try:
        import yaml
    except ImportError:
        raise DataError("YAML needs the PyYAML package (add 'pyyaml' to requirements.txt)")
    try:
        d = yaml.safe_load(text)
    except yaml.YAMLError as e:
        raise DataError(f"invalid YAML: {e}")
    if not isinstance(d, dict):
        raise DataError("the YAML must be a mapping (units: ...)")
    if {"fluid", "units", "streams"} <= set(d):
        return [], {"flowsheet": d}, ["a complete flowsheet"]
    reqs, extras, notes = [], {}, []
    units = d.get("units", d.get("parameters"))
    if units is not None:
        if not isinstance(units, dict):
            raise DataError("'units' must map unit names to parameters")
        for un, ps in units.items():
            if not isinstance(ps, dict):
                raise DataError(f"unit '{un}': expected a mapping of parameter: value")
            for k, v in ps.items():
                if k == "composition" and isinstance(v, dict):
                    for ck, cv in v.items():
                        reqs.append(_req(str(un), COMP_PREFIX + str(ck), cv, "yaml"))
                else:
                    reqs.append(_req(str(un), str(k), v, "yaml"))
    prof = d.get("profile")
    if prof is not None:
        rows = prof.get("rows") if isinstance(prof, dict) else prof
        if not isinstance(rows, list) or not rows or not all(isinstance(r, dict) for r in rows):
            raise DataError("'profile' must be a list of rows (Time and 'Unit | parameter' keys)")
        out = []
        for r in rows:
            row = {}
            for k, v in r.items():
                try:
                    row[strip_unit(k)] = to_number(v)
                except ValueError:
                    raise DataError(f"profile: '{v}' in '{k}' is not a number")
            out.append(row)
        extras["profile_rows"] = out
    evs = d.get("events", (d.get("dynamics") or {}).get("events") if isinstance(d.get("dynamics"), dict) else None)
    if evs is not None:
        if not isinstance(evs, list) or not all(isinstance(e, dict) and "kind" in e and "t" in e for e in evs):
            raise DataError("'events' must be a list of {t, kind, target, value, ramp}")
        from . import dynamics as DY
        for e in evs:
            if e["kind"] not in DY.EVENT_KINDS:
                raise DataError(f"unknown event kind '{e['kind']}'")
        extras["events"] = [dict(e) for e in evs]
    if not (reqs or extras):
        raise DataError("nothing to import: expected 'units:' (and optionally 'profile:' or 'events:')")
    notes.append("YAML")
    return reqs, extras, notes


def read_file(filename, data, model, kinds=None):
    """Everything importable in a file.  Returns {'requests': [...], 'extras': {...}, 'notes': [...], 'sheets': [(name, kind, reason)]}.
    ``kinds`` can override the detected kind of a sheet: {sheet name: kind}."""
    ext = os.path.splitext(filename)[1].lower()
    out = {"requests": [], "extras": {}, "notes": [], "sheets": []}
    if ext in (".yaml", ".yml"):
        r, e, n = requests_from_yaml(_decode(data), model)
        out.update(requests=r, extras=e, notes=n)
        return out
    if ext == ".json":
        try:
            d = json.loads(_decode(data))
        except ValueError as e:
            raise DataError(f"invalid JSON: {e}")
        if isinstance(d, dict) and {"fluid", "units", "streams"} <= set(d):
            out["extras"]["flowsheet"] = d
            out["notes"].append("a complete flowsheet")
            return out
        raise DataError("JSON import reads complete flowsheets only")
    tabs = read_tables(filename, data)
    for nm, df in tabs.items():
        kind = (kinds or {}).get(nm)
        det, reason = classify(df, model)
        use = None if kind == "Ignore" else (kind if kind in KINDS else det)
        out["sheets"].append((nm, use, reason if use == det else "chosen by you"))
        if kind == "Ignore":
            continue
        if use is None:
            out["notes"].append(f"sheet '{nm}' skipped: {reason}")
            continue
        try:
            r, e, n = requests_from_frame(df, model, use)
        except DataError as ex:
            raise DataError(f"sheet '{nm}': {ex}")
        out["requests"] += r
        for k, v in e.items():
            out["extras"].setdefault(k, v)
    if not (out["requests"] or out["extras"]):
        raise DataError("nothing importable found in the file")
    return out


# ================================================================================================== planning
def _resolve_param(u, name):
    """(key, spec or None, error) for a parameter name given for unit ``u`` (key, label, or 'label [unit]')."""
    specs = param_specs(u["type"])
    if name in specs:
        return name, specs[name], None
    low = _norm(name)
    for k, p in specs.items():
        if _norm(k) == low or _norm(p["label"]) == low or _norm(strip_unit(p["label"])) == low:
            return k, p, None
    s = strip_unit(name)
    for k, p in specs.items():
        if _norm(k) == _norm(s) or _norm(p["label"]) == _norm(s):
            return k, p, None
    if name in u["params"]:
        return name, None, None                 # a non-catalogue key that the unit already carries (e.g. a curve)
    cand = difflib.get_close_matches(name, list(specs) + [p["label"] for p in specs.values()], n=3, cutoff=0.5)
    return None, None, f"unknown parameter '{name}' for {CATALOGUE[u['type']]['label']}" + (f" (did you mean {', '.join(cand)}?)" if cand else "")


def _active(spec, params):
    """False when the parameter is hidden by the unit's current specification (show_if)."""
    si = spec.get("show_if") if spec else None
    if not si:
        return True
    for k, want in si.items():
        have = params.get(k)
        if isinstance(want, (list, tuple)):
            if have not in want:
                return False
        elif have != want:
            return False
    return True


def _coerce(spec, old, raw):
    """(value, error).  Checks type, options and min / max against the catalogue."""
    if spec is not None:
        kind = spec.get("kind")
        if kind == "float" or kind == "int":
            try:
                v = to_number(raw)
            except ValueError as e:
                return None, f"'{raw}' is not a number"
            if v is None:
                return None, "blank"
            lo, hi = spec.get("min"), spec.get("max")
            if lo is not None and v < lo - 1e-12:
                return None, f"{v:g} is below the minimum {lo:g}"
            if hi is not None and v > hi + 1e-12:
                return None, f"{v:g} is above the maximum {hi:g}"
            return v, None
        if kind == "select":
            opts = list(spec["options"])
            s = str(raw).strip()
            for o in opts:
                if str(o) == s:
                    return o, None
            for o in opts:
                if _norm(o) == _norm(s):
                    return o, None
            return None, f"'{s}' is not one of: {', '.join(map(str, opts))}"
        return str(raw), None
    # no catalogue entry: keep the type of the existing value
    if isinstance(old, bool):
        if isinstance(raw, bool):
            return raw, None
        s = _norm(raw)
        if s in ("true", "yes", "1", "on"):
            return True, None
        if s in ("false", "no", "0", "off"):
            return False, None
        return None, f"'{raw}' is not true / false"
    if isinstance(old, (int, float)):
        try:
            v = to_number(raw)
        except ValueError:
            return None, f"'{raw}' is not a number"
        return (v if v is not None else None), (None if v is not None else "blank")
    if isinstance(old, str):
        return str(raw), None
    if isinstance(raw, (dict, list)) and isinstance(old, type(raw)):
        return copy.deepcopy(raw), None
    return None, "structured parameter: set it in YAML (same shape as the current value) or in the property view"


def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-12 * max(1.0, abs(a), abs(b))
    return a == b


def plan(model, requests):
    """Check requests against the flowsheet.  One dict per change:
    {unit, uid, param, key, old, new, status, message, row}; status is ok / unchanged / error / warning / skipped.
    Later requests for the same unit and parameter override earlier ones (a note is added)."""
    if len(requests) > MAX_CHANGES:
        raise DataError(f"too many changes ({len(requests)}; the limit is {MAX_CHANGES})")
    by, dup = unit_names(model)
    comps = set(components(model))
    out, seen = [], {}
    pending = {}                                     # (uid, key) -> new value, for show_if checks of later rows
    for rq in requests:
        row = {"unit": rq["unit"], "uid": None, "param": rq["param"], "key": None, "old": None, "new": None,
               "status": "error", "message": "", "row": rq.get("row")}
        un = rq["unit"]
        if un not in by:
            cand = difflib.get_close_matches(un, list(by), n=2, cutoff=0.6)
            row["message"] = f"no unit called '{un}'" + (f" (did you mean {', '.join(cand)}?)" if cand else "")
            out.append(row)
            continue
        if un in dup:
            row["message"] = f"several units are called '{un}' - rename them first"
            out.append(row)
            continue
        uid = by[un]
        u = model["units"][uid]
        row["uid"] = uid
        pname = rq["param"]
        if pname.startswith(COMP_PREFIX):
            comp = pname[len(COMP_PREFIX):]
            if u["type"] != "feed":
                row["message"] = "compositions can only be set on feeds"
            elif comp not in comps:
                cand = difflib.get_close_matches(comp, list(comps), n=2, cutoff=0.6)
                row["message"] = f"component '{comp}' is not in the fluid package" + (f" (did you mean {', '.join(cand)}?)" if cand else "")
            else:
                try:
                    v = to_number(rq["value"])
                except ValueError:
                    v = None
                    row["message"] = f"'{rq['value']}' is not a number"
                if v is None and not row["message"]:
                    row["status"], row["message"] = "skipped", "blank"
                elif v is not None:
                    old = float((u["params"].get("composition") or {}).get(comp, 0.0) or 0.0)
                    row.update(key=pname, old=old, new=v)
                    if v < 0:
                        row.update(status="error", message="a fraction cannot be negative")
                    else:
                        row["status"] = "unchanged" if _same(old, v) else "ok"
            out.append(row)
            continue
        key, spec, err = _resolve_param(u, pname)
        if err:
            row["message"] = err
            out.append(row)
            continue
        row["key"] = key
        old = u["params"].get(key, spec.get("default") if spec else None)
        row["old"] = old
        if _is_blank(rq["value"]) and not isinstance(rq["value"], (dict, list)):
            row.update(status="skipped", message="blank")
            out.append(row)
            continue
        v, err = _coerce(spec, old, rq["value"])
        if err:
            row["message"] = err
            out.append(row)
            continue
        row["new"] = v
        row["status"] = "unchanged" if _same(old, v) else "ok"
        k2 = (uid, key)
        if k2 in seen:
            row["message"] = f"replaces the value given in row {seen[k2]}" if seen[k2] else "replaces an earlier value"
            if row["status"] == "ok":
                row["status"] = "warning"
        seen[k2] = rq.get("row")
        pending[k2] = v
        out.append(row)
    # inactive parameters (hidden by the unit's specification after all changes)
    for row in out:
        if row["status"] in ("ok", "warning") and row["uid"] and row["key"] and not row["key"].startswith(COMP_PREFIX):
            u = model["units"][row["uid"]]
            spec = param_specs(u["type"]).get(row["key"])
            if spec and spec.get("show_if"):
                eff = dict(u["params"])
                eff.update({k: v for (uid_, k), v in pending.items() if uid_ == row["uid"]})
                if not _active(spec, eff):
                    row["message"] = (row["message"] + "; " if row["message"] else "") + "not used with the unit's current specification"
                    row["status"] = "warning"
    return out


def plan_summary(pl):
    c = {s: 0 for s in STATUS_ORDER}
    for r in pl:
        c[r["status"]] = c.get(r["status"], 0) + 1
    return c


def plan_frame(pl):
    return pd.DataFrame([{"Unit": r["unit"], "Parameter": r["key"] or r["param"], "Old": _show(r["old"]), "New": _show(r["new"]),
                          "Status": r["status"], "Message": r["message"], "Row": r["row"]} for r in pl],
                        columns=["Unit", "Parameter", "Old", "New", "Status", "Message", "Row"])


def _show(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v)[:60]
    if v is None:
        return ""
    return v


def apply_plan(model, pl, accept=("ok", "warning")):
    """Write the accepted rows into the model; returns the number of values changed."""
    n = 0
    for r in pl:
        if r["status"] not in accept or r["uid"] is None:
            continue
        u = model["units"][r["uid"]]
        if r["key"].startswith(COMP_PREFIX):
            u["params"].setdefault("composition", {})[r["key"][len(COMP_PREFIX):]] = float(r["new"])
        else:
            u["params"][r["key"]] = copy.deepcopy(r["new"])
        n += 1
    return n


# ================================================================================================== batch edits
OPS = ("Set to", "Multiply by", "Add", "Reset to default")


def types_in_model(model):
    return sorted({u["type"] for u in model["units"].values()}, key=lambda t: CATALOGUE[t]["label"])


def params_of_types(types):
    """{key: label} of the parameters that occur in any of the given unit types."""
    out = {}
    for t in types:
        for p in CATALOGUE[t]["params"]:
            out.setdefault(p["key"], p["label"])
    return out


def select_units(model, types=None, pattern=""):
    """Unit ids whose type is in ``types`` (all when empty) and whose name matches ``pattern`` (wildcards * and ?,
    several patterns separated by ';' or ','; case-insensitive; no wildcard = contains)."""
    pats = [p.strip().lower() for p in re.split(r"[;,]", pattern or "") if p.strip()]
    out = []
    for uid, u in model["units"].items():
        if types and u["type"] not in types:
            continue
        nm = u["name"].lower()
        if pats and not any(fnmatch.fnmatchcase(nm, p) if any(ch in p for ch in "*?[") else p in nm for p in pats):
            continue
        out.append(uid)
    return out


def batch_requests(model, uids, key, op, value=None):
    """Requests applying one operation to one parameter of the given units (units without the parameter come back as
    'skipped' rows of the plan)."""
    if op not in OPS:
        raise DataError(f"unknown operation '{op}'")
    reqs = []
    for uid in uids:
        u = model["units"][uid]
        specs = param_specs(u["type"])
        if key.startswith(COMP_PREFIX):
            old = float((u["params"].get("composition") or {}).get(key[len(COMP_PREFIX):], 0.0) or 0.0)
            spec = {"kind": "float"}
        elif key in specs:
            spec = specs[key]
            old = u["params"].get(key, spec.get("default"))
        else:
            reqs.append({"unit": u["name"], "param": key, "value": None, "src": "batch", "row": None, "missing": True})
            continue
        if op == "Set to":
            new = value
        elif op == "Reset to default":
            new = specs[key].get("default") if key in specs else 0.0
        else:
            if spec.get("kind") not in ("float", "int"):
                reqs.append({"unit": u["name"], "param": key, "value": "not numeric", "src": "batch", "row": None})
                continue
            try:
                f = to_number(value)
            except ValueError:
                raise DataError(f"'{value}' is not a number")
            if f is None:
                raise DataError("a number is needed for this operation")
            new = float(old) * f if op == "Multiply by" else float(old) + f
        reqs.append(_req(u["name"], key, new, "batch"))
    return reqs


def plan_batch(model, uids, key, op, value=None):
    reqs = batch_requests(model, uids, key, op, value)
    miss = [r for r in reqs if r.get("missing")]
    real = [r for r in reqs if not r.get("missing")]
    pl = plan(model, real)
    for r in miss:
        pl.append({"unit": r["unit"], "uid": None, "param": key, "key": key, "old": None, "new": None, "status": "skipped",
                   "message": "this unit has no such parameter", "row": None})
    return pl


def requests_from_edited(model, before, after, key_col="Unit"):
    """Requests for the cells of a wide table (one row per unit) that differ between ``before`` and ``after``."""
    reqs = []
    cols = [c for c in after.columns if c != key_col]
    b = before.set_index(key_col)
    for _, r in after.iterrows():
        un = r[key_col]
        if _is_blank(un) or un not in b.index:
            continue
        for c in cols:
            if c not in b.columns:
                continue
            new, old = r[c], b.at[un, c]
            if _is_blank(new) and _is_blank(old):
                continue
            if not _is_blank(new) and not _is_blank(old):
                try:
                    if _same(to_number(new), to_number(old)):
                        continue
                except ValueError:
                    if str(new) == str(old):
                        continue
            reqs.append(_req(str(un), str(c), new, "table"))
    return reqs


# ================================================================================================== profile / events
def profile_settings(model, rows):
    """The Profile-tab settings that go with a table of rows: {rows, feeds, vars, extra}.  Raises DataError for a column
    that does not exist in the flowsheet."""
    from . import timeseries as TS
    try:
        bad = TS.check_columns(model, rows)
    except TS.ProfileError as e:
        raise DataError(str(e))
    if bad:
        raise DataError("; ".join(bad))
    cols = [c for c in rows[0] if c != "Time"]
    names = [n for _, n, _ in TS.feeds(model)]
    feeds = sorted({c.partition(TS.SEP)[0] for c in cols if c.partition(TS.SEP)[0] in names}) or names
    vars_ = [v for v in ("flow", "P", "T") if any(c.endswith(TS.SEP + TS.FEED_VARS[v]) for c in cols)] or ["flow", "P", "T"]
    feed_cols = set(TS.feed_columns(model, ("flow", "P", "T"), names))
    extra = [c for c in cols if c not in feed_cols]
    return {"rows": rows, "feeds": feeds, "vars": vars_, "extra": extra}


def apply_events(model, events, replace=True):
    d = model.setdefault("dynamics", {})
    d["events"] = list(events) if replace else list(d.get("events") or []) + list(events)
    return len(events)


# ================================================================================================== YAML / JSON of the flowsheet
def _plain(obj):
    return json.loads(json.dumps(obj, default=float))


def model_to_yaml(model):
    import yaml
    return yaml.safe_dump(_plain(model), sort_keys=False, allow_unicode=True, default_flow_style=False, width=120)


def parameters_to_yaml(model):
    """Editable summary: {units: {Name: {parameter: value}}} with the feed compositions - importable again."""
    import yaml
    units = OrderedDict()
    for u in model["units"].values():
        d = OrderedDict()
        for p in CATALOGUE[u["type"]]["params"]:
            d[p["key"]] = u["params"].get(p["key"], p.get("default"))
        if u["type"] == "feed":
            d["composition"] = {k: float(v) for k, v in (u["params"].get("composition") or {}).items()}
        units[u["name"]] = d
    return yaml.safe_dump({"units": _plain(units)}, sort_keys=False, allow_unicode=True, default_flow_style=False, width=120)


def check_flowsheet(d):
    """Raise DataError unless ``d`` looks like a flowsheet; return it normalised."""
    if not isinstance(d, dict) or not {"fluid", "units", "streams"} <= set(d):
        raise DataError("not a ProcessFlow Studio flowsheet (needs fluid, units and streams)")
    for uid, u in d["units"].items():
        if u.get("type") not in CATALOGUE:
            raise DataError(f"unit '{u.get('name', uid)}' has an unknown type '{u.get('type')}'")
        u.setdefault("params", {})
        u.setdefault("x", 100)
        u.setdefault("y", 100)
    d.setdefault("counter", len(d["units"]) + len(d["streams"]) + 1)
    return d


# ================================================================================================== result tables
def _num_or_text(v):
    if isinstance(v, (np.floating, np.integer)):
        return float(v)
    return v


def result_tables(model, sol=None, profile=None, dynamic=None):
    """OrderedDict of DataFrames for export.  Parameter tables are always there; result tables need a solution.
    ``profile`` is a Profile-tab result, ``dynamic`` a Dynamic-tab result dict."""
    t = OrderedDict()
    t["Parameters"] = parameter_table(model)
    t["Feed compositions"] = composition_table(model)
    if sol is not None:
        from .streams import stream_properties
        cols = OrderedDict()
        zs = OrderedDict()
        for sid, s in model["streams"].items():
            st_ = sol.streams.get(sid)
            if st_ is None:
                continue
            try:
                cols[s["name"]] = stream_properties(st_, sol.fp)
            except Exception:
                continue
            if st_.z is not None:
                zs[s["name"]] = pd.Series(np.asarray(st_.z, float), index=sol.fp.keys)
        if cols:
            t["Streams"] = pd.DataFrame(cols).rename_axis("Property").reset_index()
        if zs:
            t["Stream compositions (mole)"] = pd.DataFrame(zs).rename_axis("Component").reset_index()
        rows = []
        for uid, u in model["units"].items():
            res = sol.results.get(uid) or {}
            for k, v in res.items():
                if isinstance(v, (int, float, np.floating, np.integer, str)) and not isinstance(v, bool):
                    rows.append({"Unit": u["name"], "Type": CATALOGUE[u["type"]]["label"], "Status": sol.status.get(uid, ""),
                                 "Quantity": k, "Value": _num_or_text(v)})
        if rows:
            t["Unit results"] = pd.DataFrame(rows)
        if sol.energy:
            t["Energy streams"] = pd.DataFrame([{"Energy stream": e.name, "Unit": e.unit, "Kind": e.kind, "Duty [kW]": e.duty_kW}
                                                for e in sol.energy])
        try:
            from . import scenarios as SC
            k = SC.scenario_kpis(model, sol)
            t["Overall"] = pd.DataFrame([{"Quantity": a, "Value": _num_or_text(b)} for a, b in k.items()])
        except Exception:
            pass
        msgs = [{"Unit": model["units"][u]["name"] if u in model["units"] else u, "Status": sol.status.get(u, ""), "Message": m}
                for u, m in sol.errors.items()]
        if msgs:
            t["Messages"] = pd.DataFrame(msgs)
    if profile:
        from . import timeseries as TS
        try:
            labs = TS.labels(profile)
            t["Profile results"] = pd.DataFrame(TS.frame_rows(profile, labs))
        except Exception:
            pass
    if dynamic:
        S = dynamic["series"]
        df = pd.DataFrame({"Time [s]": dynamic["t"], **{k: v for k, v in S.items()}})
        t["Dynamic results"] = df
    ev = events_table(model)
    if len(ev):
        t["Dynamic events"] = ev
    return t


# ================================================================================================== export files
_BAD_SHEET = re.compile(r"[\[\]\:\*\?\/\\]")


def sheet_names(names):
    out, used = [], set()
    for n in names:
        s = _BAD_SHEET.sub("-", str(n)).strip("' ")[:31] or "Sheet"
        base, k = s, 2
        while s.lower() in used:
            suf = f" ({k})"
            s = base[: 31 - len(suf)] + suf
            k += 1
        used.add(s.lower())
        out.append(s)
    return out


def to_excel_bytes(tables, title=None):
    """Workbook with one formatted sheet per table (bold header, frozen header row, column widths, number format)."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    names = sheet_names(tables)
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nm, (_, df) in zip(names, tables.items()):
            d = df.copy()
            d.columns = [str(c) for c in d.columns]
            d.to_excel(xw, sheet_name=nm, index=False)
            ws = xw.sheets[nm]
            hdr_fill = PatternFill("solid", fgColor="D9ECEE")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="243746")
                cell.fill = hdr_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for j, c in enumerate(d.columns, start=1):
                vals = [str(x) for x in d[c].head(200).tolist()]
                w = min(60, max(len(str(c)), *(len(v) for v in vals)) + 2) if vals else len(str(c)) + 2
                ws.column_dimensions[get_column_letter(j)].width = max(8, w)
                if pd.api.types.is_float_dtype(d[c]):
                    for r in range(2, len(d) + 2):
                        ws.cell(row=r, column=j).number_format = "General"
        if title:
            xw.book.properties.title = str(title)[:200]
            xw.book.properties.creator = "ProcessFlow Studio"
    return buf.getvalue()


def to_csv_text(df):
    return df.to_csv(index=False)


def to_csv_zip(tables):
    buf = io.BytesIO()
    names = sheet_names(tables)
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for nm, (_, df) in zip(names, tables.items()):
            z.writestr(f"{nm}.csv", df.to_csv(index=False))
    return buf.getvalue()


# ================================================================================================== restricted Python runner
ALLOWED_IMPORTS = {"math", "statistics", "numpy", "pandas", "re", "datetime", "json", "itertools", "collections", "functools",
                   "decimal", "fractions", "textwrap", "random"}
BANNED_NAMES = {"open", "exec", "eval", "compile", "input", "getattr", "setattr", "delattr", "globals", "locals", "vars",
                "breakpoint", "memoryview", "help", "dir", "type", "super", "object", "classmethod", "staticmethod", "property",
                "exit", "quit", "__import__", "__builtins__"}
BANNED_ATTRS = {"to_csv", "to_excel", "to_json", "to_pickle", "to_parquet", "to_sql", "to_hdf", "to_feather", "to_orc",
                "to_stata", "to_xml", "to_clipboard", "to_latex", "to_markdown", "to_gbq", "eval", "load", "loadtxt",
                "genfromtxt", "fromfile", "save", "savez", "savez_compressed", "savetxt", "tofile", "memmap", "ctypeslib",
                "f2py", "testing", "system", "popen", "fromregex", "DataSource", "lib", "distutils", "show_config", "read_clipboard",
                # module / frame / file-system doors found in the security review
                "io", "os", "sys", "core", "compat", "util", "ctypes", "builtins", "modules", "subprocess", "importlib", "pickle",
                "style", "query", "plot", "hist", "boxplot", "ExcelFile", "ExcelWriter", "HDFStore", "get_include", "show_runtime",
                "who", "lookfor", "gi_frame", "gi_code", "cr_frame", "ag_frame", "f_globals", "f_builtins",
                "f_locals", "f_back", "tb_frame", "tb_next", "func_globals", "co_code", "with_traceback", "mro", "format_map",
                "pipe_to", "shutil", "socket", "signal", "inspect", "Formatter", "Template"}
# attributes that start with to_ and are fine (everything else starting with to_ writes somewhere)
TO_OK = {"to_numeric", "to_datetime", "to_timedelta", "to_dict", "to_list", "to_numpy", "to_frame", "to_series", "to_period",
         "to_timestamp", "to_records", "to_string", "to_flat_index", "to_pydatetime", "to_pytimedelta", "to_julian_date",
         "to_perioddelta", "to_native_types", "to_decimal", "to_integral_value", "to_integral_exact", "to_eng_string"}
# sub-modules the script may reach through numpy / pandas (everything else that is a module is hidden)
SUBMODULES_OK = {"random", "linalg", "fft", "polynomial", "api", "types"}
MAX_CODE = 40000
MAX_CELLS = 3_000_000


def check_code(code):
    """Raise DataError for code that tries to import other modules, touch files or reach into Python internals.
    First line of defence only: the child process also hides module attributes behind proxies, audits file / process /
    network calls and runs with resource limits.  It is still not a hardened sandbox - see docs/VALIDATION.md."""
    if len(code) > MAX_CODE:
        raise DataError(f"the script is longer than {MAX_CODE} characters")
    try:
        tree = ast.parse(code)
    except SyntaxError as e:
        raise DataError(f"syntax error on line {e.lineno}: {e.msg}")
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for a in node.names:
                if a.name.split(".")[0] not in ALLOWED_IMPORTS or "." in a.name:
                    raise DataError(f"line {node.lineno}: 'import {a.name}' is not allowed (allowed: {', '.join(sorted(ALLOWED_IMPORTS))}; no submodules)")
        elif isinstance(node, ast.ImportFrom):
            if node.level or (node.module or "").split(".")[0] not in ALLOWED_IMPORTS or "." in (node.module or ""):
                raise DataError(f"line {node.lineno}: 'from {node.module} import ...' is not allowed (allowed: {', '.join(sorted(ALLOWED_IMPORTS))})")
            for a in node.names:
                if a.name.startswith("_") or a.name in BANNED_ATTRS or a.name.startswith("read_"):
                    raise DataError(f"line {node.lineno}: '{a.name}' is not available in the editor")
        elif isinstance(node, ast.Name):
            if node.id in BANNED_NAMES or node.id.startswith("__"):
                raise DataError(f"line {node.lineno}: '{node.id}' is not available in the editor")
        elif isinstance(node, ast.Attribute):
            a = node.attr
            if a.startswith("_") or a in BANNED_ATTRS or a.startswith("read_") or (a.startswith("to_") and a not in TO_OK):
                raise DataError(f"line {node.lineno}: '.{a}' is not available in the editor (no file access, no private or internal attributes)")
        elif isinstance(node, ast.Constant) and isinstance(node.value, str) and "__" in node.value:
            raise DataError(f"line {node.lineno}: text containing '__' is not allowed")
        elif isinstance(node, (ast.Global, ast.Nonlocal)):
            raise DataError(f"line {node.lineno}: global / nonlocal is not available")
    return tree


_RUNNER = r'''
import sys, io, json, math, time, builtins, contextlib, traceback
try:
    import resource
    cpu = int(sys.argv[1])
    resource.setrlimit(resource.RLIMIT_CPU, (cpu, cpu + 1))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    try:
        resource.setrlimit(resource.RLIMIT_AS, (3 * 1024 ** 3, 3 * 1024 ** 3))
    except Exception:
        pass
except Exception:
    pass
import numpy as np, pandas as pd
payload = json.loads(sys.stdin.read())
code = payload["code"]
tables = {k: pd.read_json(io.StringIO(v), orient="split", precise_float=True) for k, v in payload["tables"].items()}
ALLOWED = set(payload["allowed"])
BANNED = set(payload["banned"])
SUBOK = set(payload["submodules"])
import types as _types, os as _os, site as _site
_real_import = builtins.__import__

class _Mod:
    """Read-only view of a module: hides private names, banned names and (almost) all sub-modules, so a script cannot walk
    from pd / np to os, sys or io."""
    __slots__ = ("_m", "_n")
    def __init__(self, m, n):
        object.__setattr__(self, "_m", m)
        object.__setattr__(self, "_n", n)
    def __getattr__(self, name):
        if name.startswith("_") or name in BANNED or name.startswith("read_"):
            raise AttributeError("'%s' is not available in the editor" % name)
        v = getattr(object.__getattribute__(self, "_m"), name)
        if isinstance(v, _types.ModuleType):
            if name in SUBOK:
                return _Mod(v, object.__getattribute__(self, "_n") + "." + name)
            raise AttributeError("module '%s' is not available in the editor" % name)
        return v
    def __setattr__(self, name, value):
        raise AttributeError("modules are read-only in the editor")
    def __dir__(self):
        return [k for k in dir(object.__getattribute__(self, "_m")) if not k.startswith("_") and k not in BANNED]
    def __repr__(self):
        return "<module %s (restricted)>" % object.__getattribute__(self, "_n")

_cache = {}
def _view(name):
    if name not in _cache:
        _cache[name] = _Mod(_real_import(name), name)
    return _cache[name]

def _imp(name, globals=None, locals=None, fromlist=(), level=0):
    if level or name.split(".")[0] not in ALLOWED or "." in name:
        raise ImportError("import of '%s' is not allowed" % name)
    return _view(name)

# defence in depth: refuse file reads outside the Python installation, any file write, processes, sockets, ctypes
_ok_roots = tuple({_os.path.realpath(p) for p in [sys.prefix, sys.base_prefix, sys.exec_prefix, sys.base_exec_prefix] +
                   list(getattr(_site, "getsitepackages", lambda: [])()) + [_site.getusersitepackages()] +
                   [p for p in sys.path if p]})
_ok_roots = tuple(r for r in _ok_roots if r not in ("/", ""))
_ok_roots += ("/proc/cpuinfo", "/proc/meminfo", "/sys/devices/system/cpu", "/dev/null", "/dev/urandom", "/etc/localtime", "/usr/share/zoneinfo")
_DENY = ("subprocess.", "os.system", "os.exec", "os.spawn", "os.posix_spawn", "os.fork", "os.forkpty", "os.kill", "os.remove",
         "os.rename", "os.rmdir", "os.mkdir", "os.chmod", "os.chown", "os.truncate", "os.symlink", "os.link", "os.putenv",
         "os.chdir", "os.chroot", "os.startfile", "socket.", "ctypes.", "shutil.", "pty.", "webbrowser.", "urllib.Request",
         "http.client", "ftplib.", "smtplib.", "telnetlib.", "pickle.find_class", "sqlite3.", "winreg.")
def _audit(ev, args):
    if ev.startswith(_DENY):
        raise PermissionError("blocked by the editor: " + ev)
    if ev == "open":
        path, mode = args[0], args[1] if len(args) > 1 else "r"
        if isinstance(path, int):
            return
        if mode and any(c in str(mode) for c in "wax+"):
            raise PermissionError("the editor cannot write files")
        try:
            rp = _os.path.realpath(_os.fsdecode(path))
        except Exception:
            raise PermissionError("blocked by the editor: open")
        if not rp.startswith(_ok_roots):
            raise PermissionError("the editor cannot read files (" + _os.path.basename(rp) + ")")
sys.addaudithook(_audit)
SAFE = {k: getattr(builtins, k) for k in (
    "abs all any bool bytes dict divmod enumerate filter float format frozenset hash int isinstance issubclass iter len list map "
    "max min next pow print range repr reversed round set slice sorted str sum tuple zip callable chr ord bin hex oct "
    "Exception ValueError TypeError KeyError IndexError ZeroDivisionError ArithmeticError StopIteration RuntimeError "
    "True False None NotImplemented Ellipsis").split() if hasattr(builtins, k)}
SAFE["__import__"] = _imp
out = io.StringIO()
res = {"ok": True, "error": "", "stdout": "", "tables": {}}
ns = {"__builtins__": SAFE, "pd": _view("pandas"), "np": _view("numpy"), "math": _view("math"), "tables": tables}
try:
    with contextlib.redirect_stdout(out):
        exec(compile(code, "<editor>", "exec"), ns)
    t = ns.get("tables")
    if not isinstance(t, dict):
        raise TypeError("'tables' must stay a dict of DataFrames")
    for k, v in t.items():
        if isinstance(v, pd.Series):
            v = v.to_frame()
        if not isinstance(v, pd.DataFrame):
            raise TypeError("tables[%r] is not a DataFrame (it is %s)" % (k, type(v).__name__))
        v = v.copy()
        v.columns = [str(c) for c in v.columns]
        res["tables"][str(k)] = v.to_json(orient="split", date_format="iso", double_precision=15)
except BaseException as e:
    res["ok"] = False
    tb = traceback.extract_tb(e.__traceback__)
    line = [f.lineno for f in tb if f.filename == "<editor>"]
    res["error"] = "%s: %s%s" % (type(e).__name__, e, (" (line %d)" % line[-1]) if line else "")
res["stdout"] = out.getvalue()[:20000]
sys.stdout.write("\n@@RESULT@@" + json.dumps(res))
'''


def run_script(code, tables, timeout=30):
    """Run ``code`` on copies of ``tables`` in a separate Python process with CPU / memory / file-size limits and a
    wall-clock timeout.  The script sees ``tables`` (dict of DataFrames), ``pd``, ``np`` and ``math`` and changes the
    tables in place (or replaces / adds entries).  Returns {'ok', 'error', 'stdout', 'tables', 'seconds'}."""
    t0 = time.time()
    check_code(code)
    cells = sum(df.shape[0] * df.shape[1] for df in tables.values())
    if cells > MAX_CELLS:
        raise DataError(f"the tables are too large for the editor ({cells:,} cells; limit {MAX_CELLS:,})")
    payload = json.dumps({"code": code, "allowed": sorted(ALLOWED_IMPORTS), "banned": sorted(BANNED_ATTRS), "submodules": sorted(SUBMODULES_OK),
                          "tables": {k: df.to_json(orient="split", date_format="iso", double_precision=15) for k, df in tables.items()}})
    env = {"PATH": os.environ.get("PATH", ""), "OMP_NUM_THREADS": "1", "OPENBLAS_NUM_THREADS": "1", "PYTHONHASHSEED": "0",
           "HOME": tempfile.gettempdir(), "LANG": "C.UTF-8"}
    try:
        with tempfile.TemporaryDirectory() as td:
            p = subprocess.run([sys.executable, "-I", "-c", _RUNNER, str(int(timeout))], input=payload, capture_output=True,
                               text=True, timeout=timeout + 5, cwd=td, env=env)
    except subprocess.TimeoutExpired:
        return {"ok": False, "error": f"stopped after {timeout} s (infinite loop or a very slow calculation?)", "stdout": "",
                "tables": None, "seconds": time.time() - t0}
    txt = p.stdout
    if "@@RESULT@@" not in txt:
        err = (p.stderr or "").strip().splitlines()
        msg = err[-1] if err else f"the process ended with code {p.returncode}"
        if p.returncode in (-9, -24, -25, 137):
            msg = "stopped: the script used too much CPU time or memory"
        return {"ok": False, "error": msg, "stdout": txt[:5000], "tables": None, "seconds": time.time() - t0}
    head, _, body = txt.rpartition("@@RESULT@@")
    res = json.loads(body)
    tabs = None
    if res["ok"]:
        tabs = OrderedDict((k, pd.read_json(io.StringIO(v), orient="split", precise_float=True)) for k, v in res["tables"].items())
    return {"ok": res["ok"], "error": res["error"], "stdout": (head.strip() + "\n" + res["stdout"]).strip()[:20000],
            "tables": tabs, "seconds": time.time() - t0}


def _cell_equal(x, y):
    if _is_blank(x) and _is_blank(y):
        return True
    try:
        if isinstance(x, (int, float, np.integer, np.floating)) and isinstance(y, (int, float, np.integer, np.floating)) \
                and not isinstance(x, bool) and not isinstance(y, bool):
            return bool(np.isclose(float(x), float(y), rtol=1e-9, atol=0.0))
    except (TypeError, ValueError):
        pass
    return x == y


def diff_tables(before, after):
    """[(table, status, detail)] comparing two dicts of DataFrames: added / removed / changed / unchanged.  Numbers that
    differ by less than 1e-9 relative (float noise of the round trip) are the same."""
    out = []
    for k in after:
        if k not in before:
            out.append((k, "added", f"{after[k].shape[0]} rows x {after[k].shape[1]} columns"))
            continue
        a, b = before[k], after[k]
        if a.shape != b.shape or list(map(str, a.columns)) != list(map(str, b.columns)):
            out.append((k, "changed", f"{a.shape[0]}x{a.shape[1]} -> {b.shape[0]}x{b.shape[1]}"))
            continue
        av, bv = a.reset_index(drop=True).to_numpy(dtype=object), b.reset_index(drop=True).to_numpy(dtype=object)
        n = sum(0 if _cell_equal(x, y) else 1 for x, y in zip(av.ravel(), bv.ravel()))
        out.append((k, "changed" if n else "unchanged", f"{n} cell(s) differ" if n else ""))
    for k in before:
        if k not in after:
            out.append((k, "removed", ""))
    return out


SNIPPETS = OrderedDict([
    ("Round every number to 2 decimals", "for name, df in tables.items():\n    num = df.select_dtypes('number').columns\n    df[num] = df[num].round(2)\n"),
    ("Keep only the rows of one unit", "df = tables['Unit results']\ntables['Unit results'] = df[df['Unit'].str.contains('K-100')]\n"),
    ("Convert stream pressures from bar to kPa (new row)",
     "s = tables['Streams']\nrow = s[s['Property'] == 'Pressure [bar(a)]'].copy()\nrow['Property'] = 'Pressure [kPa(a)]'\nfor c in row.columns[1:]:\n    row[c] = row[c].astype(float) * 100\ntables['Streams'] = pd.concat([s, row], ignore_index=True)\n"),
    ("Add a summary sheet", "df = tables['Unit results']\nnum = df[df['Quantity'].str.contains('Power')]\ntables['Summary'] = pd.DataFrame({'Total power [kW]': [pd.to_numeric(num['Value'], errors='coerce').sum()]})\n"),
    ("Scale a parameter on all units of a name pattern (for re-import)",
     "p = tables['Parameters']\nm = p['Unit'].str.contains('Cooler') & (p['Parameter'] == 'dP')\np.loc[m, 'Value'] = p.loc[m, 'Value'].astype(float) * 0.9\n"),
    ("Drop a table from the export", "tables.pop('Messages', None)\n"),
])
